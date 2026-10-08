#!/usr/bin/env python3
"""원본 발표 덱을 건드리지 않고 _v2.pptx 와 대본 .md 를 만든다.

    .venv-cad/Scripts/python assets/cocktail_rh56/presentation/build_deck_v2.py
"""
import copy
import io
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from script_20261002 import SCRIPT, STORYLINE  # noqa: E402

SRC = r"C:\Users\tony0\Downloads\비전기반_TaskPlanner_연구탐구_20261002.pptx"
DST = SRC.replace(".pptx", "_v2.pptx")
MD = os.path.join(os.path.dirname(SRC), "발표대본_TaskPlanner_20261002.md")

U = 91440  # 0.1 inch; slide = 160 x 120 units
FONT = "Pretendard Variable"
INK, BODY, ACC, PALE, MUTE, LINE = "222222", "45515E", "1456F0", "EEF4FF", "8E8E93", "D9E1EC"


def u(v):
    return Emu(int(v * U))


def rgb(h):
    return RGBColor.from_string(h)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, gap=4):
    """paras: list of (text, size, bold, color) or list of such lists for runs."""
    tb = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(4)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        runs = para if isinstance(para, list) else [para]
        for t, size, bold, color in runs:
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
            r.font.color.rgb = rgb(color)
    return tb


def box(slide, x, y, w, h, fill="FFFFFF", line=LINE, rounded=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               u(x), u(y), u(w), u(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    if rounded:
        s.adjustments[0] = 0.06
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def set_first_run(shape, new):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def new_slide_like(prs, tmpl, title):
    s = prs.slides.add_slide(tmpl.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    for nm in ("Text 0", "Shape 1", "Text 3"):
        s.shapes._spTree.append(copy.deepcopy(shape(tmpl, nm)._element))
    for nm in ("Image 0", "Image 1"):
        im = shape(tmpl, nm)
        s.shapes.add_picture(io.BytesIO(im.image.blob), im.left, im.top, im.width, im.height)
    set_first_run(shape(s, "Text 0"), title)
    return s


def move(prs, slide, index):
    lst = prs.slides._sldIdLst
    el = next(e for e in lst if int(e.get("id")) == slide.slide_id)
    lst.remove(el)
    lst.insert(index, el)


# ---------------------------------------------------------------------------
prs = Presentation(SRC)
S = list(prs.slides)
s_contents, s_old8, s_old9, s_old10 = S[1], S[7], S[8], S[9]

# --- 목차: 3·4 항목을 실제 내용에 맞춘다
set_first_run(shape(s_contents, "Text 18"), "내 파트 — RH56 한 손 쉐이킹을 위한 중력 잠금 셰이커 리드 설계")
set_first_run(shape(s_contents, "Text 23"), "Sim 학습 · 실물 검증 · Planner Skill/Recovery 통합")

# --- 기존 3장 제목 구체화 + 캡션
set_first_run(shape(s_old8, "Text 0"), "3. 진행상황 · 설계 1단계: 그립존 축소 → 한 손 쉐이킹 + 홀더")
set_first_run(shape(s_old9, "Text 0"), "3. 진행상황 · 설계 2단계: 중력 리드인 + 바요넷 잠금 + 홀더 키")
set_first_run(shape(s_old10, "Text 0"), "3. 진행상황 · 결합 과정: 놓고 → 비틀고 → 잠김")
cap = lambda s, x, y, w, t, sub: text(s, x, y, w, 12, [(t, 13, True, INK), (sub, 11, False, BODY)])  # noqa: E731
cap(s_old9, 4, 40, 31, "분해도", "홀더 · 바디 · 리드 3파트")
cap(s_old9, 127, 26, 31, "잠긴 상태 (리드 반투명)", "러그 4개가 4 mm 잠금 턱 아래")
cap(s_old9, 119, 82, 38, "홀더 키 6개", "바디 회전·딸려 올라옴 방지 · ±30° 허용")
text(s_old10, 78, 55.5, 70, 5, [[("결합 애니메이션  ", 12, True, INK),
                                 ("+28° 어긋나게 놓아도 자동 정렬 → 30° 비틀어 잠금", 11, False, BODY)]])
box(s_old10, 146, 60, 13, 50, fill=PALE, line=None)
text(s_old10, 146.5, 61, 12, 48, [("로봇 허용 오차", 11, True, ACC), ("회전", 10, False, BODY), ("±45°", 16, True, INK),
                                    ("위치", 10, False, BODY), ("4.5 mm", 16, True, INK),
                                    ("경사 40°", 10, False, BODY), ("μ < 0.84면", 10, False, BODY),
                                    ("중력만으로 정렬", 10, False, BODY)])

# --- 새 슬라이드 A: 내 파트와 제약 (기존 8번 앞)
A = new_slide_like(prs, s_old8, "3. 진행상황 · 내 파트: RH56으로 셰이커 리드를 결합하고 흔들기")
box(A, 6, 16, 70, 64)
text(A, 9, 18, 64, 5, [("TASK (내 담당)", 13, True, ACC)])
text(A, 9, 24, 64, 54, [
    ("시작 상태", 18, True, INK),
    ("앞 작업자 Skill로 컵·리드를 이미 파지한 상태", 16, False, BODY),
    ("목표", 18, True, INK),
    ("리드를 peg-in-hole처럼 정확히 결합 → 한 손 쉐이킹 → 홀더에 놓고 리드 분리", 16, False, BODY),
    ("Planner 관점", 18, True, INK),
    ("Skill Library의 'Lid-on / Shake' Skill 하나. 성공 여부가 보여야 Recovery가 동작한다", 16, False, BODY),
], gap=10)
box(A, 82, 16, 72, 64)
text(A, 85, 18, 66, 5, [("RH56F1 제약 (URDF 순기구학 추정)", 13, True, ACC)])
rows = [("RH56 엄지–손바닥 최대 거리", "약 70–80 mm"), ("안정적 감싸쥐기 직경", "약 50–65 mm"),
        ("기존 Cobbler 셰이커 최대 외경", "92 mm"), ("기존 RL 태스크의 축소 셰이커", "57 mm (×0.65)"),
        ("새 설계 그립 외경", "60 mm")]
for i, (k, v) in enumerate(rows):
    y = 25 + i * 10
    box(A, 85, y, 66, 8.5, fill=PALE if i in (2, 4) else "FFFFFF", line=LINE, rounded=False)
    text(A, 86, y, 42, 8.5, [(k, 15, False, BODY)], anchor=MSO_ANCHOR.MIDDLE)
    text(A, 128, y, 22, 8.5, [(v, 17, True, ACC if i == 4 else INK)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
text(A, 85, 76, 66, 4, [("* 메시 기반 추정치. 실린더 직경 스윕 시뮬로 확정 예정", 10, False, MUTE)])
box(A, 6, 85, 148, 22, fill=ACC, line=None)
text(A, 9, 86, 142, 20, [
    ("두 손 쉐이킹은 RL로 풀기 어렵다 → 한 손 쉐이킹 → 리드가 스스로 단단히 고정돼야 한다", 20, True, "FFFFFF"),
    ("그래서 정책을 키우는 대신, 물체(셰이커·리드·홀더)를 다시 설계해 Skill을 쉽게 만든다", 16, False, "DBE7FF"),
], anchor=MSO_ANCHOR.MIDDLE, gap=8)

# --- 새 슬라이드 B: Future Work (맨 끝)
B = new_slide_like(prs, s_old8, "4. Future Work: Sim 학습 → 실물 검증 → Planner 통합")
cols = [
    ("① 시뮬레이션", ["Isaac Lab에서 drop → twist 정책 학습",
                     "sim 프로파일(간극 1.0 mm) USD 변환 · SDF 충돌",
                     "성공 판정: 리드 yaw −30° · z −0.9 mm"]),
    ("② 실물 검증", ["FDM 출력 (real 프로파일, 간극 0.4 mm)",
                   "딸깍 감 · 개스킷 압착 확인",
                   "RH56 손가락으로 30° 비틀 토크가 충분한가"]),
    ("③ Planner 통합", ["'Lid-on'을 Skill Library에 등록",
                       "잠금 상태(yaw·z)를 Recovery 실패 감지 신호로",
                       "RoboDexVLM식 dual-layer recovery와 연결"]),
]
for i, (h, items) in enumerate(cols):
    x = 6 + i * 50
    box(B, x, 17, 47, 52)
    text(B, x + 3, 20, 41, 7, [(h, 20, True, ACC)])
    text(B, x + 3, 30, 41, 38, [("· " + t, 16, False, BODY) for t in items], gap=14)
box(B, 6, 73, 148, 15, fill=PALE, line=None)
text(B, 9, 74, 142, 13, [[("리스크  ", 15, True, INK),
                          ("리드 노브(폭 36 mm)를 RH56이 안정적으로 쥘 수 있는가 · 흔들 때 컵이 손에서 미끄러지지 않는가 · "
                           "현재 검사는 설계 수치 기반이며 접촉 시뮬과 실물은 아직", 14, False, BODY)]],
     anchor=MSO_ANCHOR.MIDDLE)
box(B, 6, 92, 148, 16, fill=ACC, line=None)
text(B, 9, 93, 142, 14, [("어려운 Skill은 정책만 키우지 말고 물체 설계로 쉽게 만들고, "
                          "그 결과를 Planner가 읽을 수 있는 신호로 만든다", 19, True, "FFFFFF")],
     anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

move(prs, A, 7)

# --- 페이지 번호 + 발표자 노트
for i, s in enumerate(prs.slides, 1):
    if i > 1:
        set_first_run(shape(s, "Text 3"), str(i))
body_ph = next(sh for sh in S[2].notes_slide.placeholders if sh.placeholder_format.type == 2)  # BODY
for s, (title, sec, body, tip) in zip(prs.slides, SCRIPT):
    ns = s.notes_slide
    if ns.notes_text_frame is None:  # 이 덱의 notes master 에는 본문 placeholder 가 없어 새 슬라이드엔 직접 복사
        ns.shapes._spTree.append(copy.deepcopy(body_ph._element))
    tf = ns.notes_text_frame
    tf.text = f"({sec}초) {body}" + (f"\n\n[진행 팁] {tip}" if tip else "")

prs.save(DST)

# --- 대본 md
CPS = 5.0  # 한국어 발표 속도 가정: 초당 약 5자(공백 포함)
lines = ["# 비전 기반 Task Planner 연구 탐구 — 발표 대본 (2026.10.02)", "",
         f"- 대상 파일: `{os.path.basename(DST)}` (원본은 수정하지 않음)",
         "- 발표 10분 + 질의응답 5분. 시간은 초당 약 5자(공백 포함)로 말한다고 가정한 추정치입니다.", "",
         "## 스토리라인", ""]
lines += [f"{i}. **{k}** — {v}" for i, (k, v) in enumerate(STORYLINE, 1)]
lines += ["", "## 시간 배분", "", "| # | 슬라이드 | 목표 | 대본 분량 기준 |", "|---|---|---|---|"]
tot_t = tot_e = 0
for i, (title, sec, body, tip) in enumerate(SCRIPT, 1):
    est = len(body) / CPS
    tot_t, tot_e = tot_t + sec, tot_e + est
    lines.append(f"| {i} | {title} | {sec}초 | {est:.0f}초 |")
lines.append(f"| | **합계** | **{tot_t // 60}분 {tot_t % 60}초** | **{int(tot_e) // 60}분 {int(tot_e) % 60}초** |")
lines += ["", "## 대본", ""]
for i, (title, sec, body, tip) in enumerate(SCRIPT, 1):
    lines += [f"### {i}. {title}  ({sec}초)", "", body, ""]
    if tip:
        lines += [f"> 진행 팁: {tip}", ""]
with open(MD, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(DST)
print(MD)
print(f"target {tot_t}s  script-length estimate {tot_e:.0f}s  slides {len(prs.slides)}")
